import math
import os
import re
import json
import base64
import io
import hashlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi import FastAPI, HTTPException, Form, Query, UploadFile, File, Request, Depends
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from dateutil.relativedelta import relativedelta
from pydantic import BaseModel
import qrcode
from PIL import Image, ImageDraw, ImageFont
from datetime import datetime, date
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded

from database import banco_dados
from auth import get_current_user, require_admin, hash_senha, verificar_senha, criar_token

# uvicorn main:app --reload

app = FastAPI()

limiter = Limiter(key_func=get_remote_address)
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)

FRONTEND_ORIGINS = [
    origem.strip()
    for origem in os.getenv(
        "FRONTEND_ORIGINS",
        "http://localhost:3000,https://qr-code-hunt.vercel.app",
    ).split(",")
    if origem.strip()
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=FRONTEND_ORIGINS,
    allow_methods=["*"],
    allow_headers=["*"],
)

ENUMS_PERMITIDOS = {"escola", "curso_interesse",  "status_academico"}
PONTOS_POR_AMIGO = 50

@app.get("/opcoes/{nome}")
def opcoes(nome: str):
    """Busca os valores do enum diretamente do banco de dados."""
    if nome not in ENUMS_PERMITIDOS:
        raise HTTPException(status_code=404, detail=f"Enum '{nome}' não encontrado.")
    res = banco_dados.rpc("get_enum_values", {"enum_name": nome}).execute()
    return res.data

def validar_email_backend(email: str):
    padrao = r"^[a-z0-9!#$%&'*+/=?^_`{|}~-]+(?:\.[a-z0-9!#$%&'*+/=?^_`{|}~-]+)*@(?:[a-z0-9](?:[a-z0-9-]*[a-z0-9])?\.)+[a-z0-9](?:[a-z0-9-]*[a-z0-9])?$"
    if not re.match(padrao, email.lower()):
        return False, "E-mail inválido."
    return True, ""

def validar_nome_sem_numeros(nome: str):
    if any(char.isdigit() for char in nome):
        return False, "O nome não pode conter números."
    if len(nome.strip()) < 6:
        return False, "Nome muito curto."
    return True, ""

def validar_senha(senha: str):
    if len(senha) < 8:
        return False, "A senha precisa ter pelo menos 8 caracteres."
    return True, ""

def calcular_nivel(pontos: int) -> int:
    return int(math.sqrt((pontos or 0) / 50)) + 1

def gerar_code_hash(semente: str) -> str:
    dados_hash = f"{semente}{os.urandom(8).hex()}"
    return hashlib.sha256(dados_hash.encode()).hexdigest()[:16]

@app.get("/")
def read_index():
    return RedirectResponse(url="/docs")

# ==================== AUTH ====================

@app.get("/auth/me")
def auth_me(id_user: str = Depends(get_current_user)):
    res = banco_dados.table("users").select("id_user, nome, email, pontos, is_admin").eq("id_user", id_user).single().execute()
    if not res.data:
        raise HTTPException(status_code=404, detail="Usuário não encontrado.")
    dados = res.data
    dados["nivel"] = calcular_nivel(dados.get("pontos"))
    return dados

# ==================== USER ====================

@app.post("/usuarios/novo")
@limiter.limit("10/minute")
def cadastro_user(
    request: Request,
    nome: str = Form(...),
    email: str = Form(...),
    senha: str = Form(...),
    data_nasc: str = Form(...),
    telefone: str = Form(""),
    escola: str = Form(""),
):
    v_nome, m_nome = validar_nome_sem_numeros(nome)
    if not v_nome: raise HTTPException(status_code=400, detail=m_nome)

    v_email, m_email = validar_email_backend(email)
    if not v_email: raise HTTPException(status_code=400, detail=m_email)

    v_senha, m_senha = validar_senha(senha)
    if not v_senha: raise HTTPException(status_code=400, detail=m_senha)

    data_n_dt = datetime.strptime(data_nasc, "%Y-%m-%d")
    idade = relativedelta(datetime.now(), data_n_dt).years
    if idade < 15:
        raise HTTPException(status_code=400, detail="Você precisa ter pelo menos 15 anos.")


    if not escola.strip():
        raise HTTPException(status_code=400, detail="Selecione uma instituição de ensino.")

    registro = datetime.now().isoformat()
    dados_user = {
        "nome": nome.title(),
        "data_nasc": data_nasc,
        "email": email.lower().strip(),
        "senha_hash": hash_senha(senha),
        "telefone": telefone,
        "escola": escola or None,
        "pontos": 0,
        "data_registro": registro
    }

    try:
        response = banco_dados.table("users").insert(dados_user).execute()
        return {"status": "Sucesso", "user": response.data[0]}
    except Exception as e:
        print(f"Erro detectado: {e}")
        error_msg = str(e)
        if "23505" in error_msg or "duplicate key" in error_msg:
            raise HTTPException(status_code=400, detail="Este e-mail já está cadastrado em nossa base.")
        raise HTTPException(status_code=500, detail="Erro interno no servidor ao realizar cadastro.")

@app.get("/usuarios/dados/exportar")
def exportar_dados(
    admin_id: str = Depends(require_admin),
    data: date = Query(...),
    formato: str = Query("xlsx"),
    pontos_min: int = Query(None),
    pontos_max: int = Query(None),
):
    query = (
        banco_dados.table("users")
        .select("nome", "pontos", "escola", 'curso_interesse', "status_academico")
        .gte("data_registro", f"{data}T00:00:00")
        .lte("data_registro", f"{data}T23:59:59")
        .neq("email", "softwarehouseundb@gmail.com")
    )
    if pontos_min is not None:
        query = query.gte("pontos", pontos_min)
    if pontos_max is not None:
        query = query.lte("pontos", pontos_max)

    response = query.execute()
    if not response.data:
        raise HTTPException(status_code=404, detail=f"Nenhum usuário encontrado para os filtros informados.")

    usuarios = response.data

    if formato == "json":
        relatorio = {
            "data_relatorio": str(data),
            "filtros": { "pontos_min": pontos_min, "pontos_max": pontos_max},
            "total_alunos": len(usuarios),
            "alunos": usuarios
        }
        json_bytes = json.dumps(relatorio, ensure_ascii=False, indent=2).encode("utf-8")
        return StreamingResponse(io.BytesIO(json_bytes), media_type="application/json",
            headers={"Content-Disposition": f"attachment; filename=relatorio_{data}.json"})

    elif formato == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Relatório {data}"
        colunas = ["Nome", "Pontos",  "Escola"]
        campos  = ["nome", "pontos", "escola"]
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for col_idx, titulo in enumerate(colunas, start=1):
            cell = ws.cell(row=1, column=col_idx, value=titulo)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row_idx, usuario in enumerate(usuarios, start=2):
            for col_idx, campo in enumerate(campos, start=1):
                ws.cell(row=row_idx, column=col_idx, value=usuario.get(campo, ""))
        for col in ws.columns:
            max_length = max((len(str(c.value)) for c in col if c.value), default=10)
            ws.column_dimensions[col[0].column_letter].width = max_length + 4
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return StreamingResponse(buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=relatorio_{data}.xlsx"})

    raise HTTPException(status_code=400, detail="Formato inválido.")

@app.get("/ranking")
def ranking(limit: int = 10):
    """
    Retorna o ranking ordenado por pontos.
    - limit=10  → top 10 para a tela de ranking/home (padrão)
    - limit=0   → todos os usuários
    """
    query = (
        banco_dados.table("users")
        .select("id_user, nome, pontos, catch(count), data_registro")
        .eq("is_admin", False)
        .order("pontos", desc=True)
        .order("data_registro", desc=False)
    )
    if limit > 0:
        query = query.limit(limit)

    res = query.execute()

    ranking_formatado = []
    for user in res.data:
        ranking_formatado.append({
            "id": user["id_user"],
            "nome": user["nome"],
            "pontos": user["pontos"],
            "nivel": calcular_nivel(user["pontos"]),
            "qrs_capturados": user["catch"][0]["count"] if user.get("catch") else 0
        })
    return ranking_formatado


@app.get("/usuarios/me/posicao")
def posicao_usuario(id_user: str = Depends(get_current_user)):
    """
    Retorna a posição, pontos e QRs do usuário autenticado sem carregar o ranking inteiro.
    Conta quantos usuários não-admin têm pontos maiores que o alvo.
    """
    usuario = (
        banco_dados.table("users")
        .select("id_user, nome, pontos, data_registro")
        .eq("id_user", id_user)
        .single()
        .execute()
    )
    if not usuario.data:
        raise HTTPException(status_code=404, detail="Usuário não encontrado.")

    pontos = usuario.data["pontos"]
    acima_por_pontos = (
        banco_dados.table("users")
        .select("id_user", count="exact")
        .eq("is_admin", False)
        .gt("pontos", pontos)
        .execute()
    )

    acima_por_registro = (
        banco_dados.table("users")
        .select("id_user", count="exact")
        .eq("is_admin", False)
        .eq("pontos", pontos)
        .lt("data_registro", usuario.data["data_registro"])
        .execute()
    )

    posicao = (acima_por_pontos.count or 0) + (acima_por_registro.count or 0) + 1

    qrs = (
        banco_dados.table("catch")
        .select("id_catch", count="exact")
        .eq("id_user", id_user)
        .execute()
    )

    return {
        "id": id_user,
        "nome": usuario.data["nome"],
        "pontos": pontos,
        "nivel": calcular_nivel(pontos),
        "posicao": posicao,
        "qrs_capturados": qrs.count or 0
    }

@app.post("/login")
@limiter.limit("5/minute")
def login(request: Request, email: str = Form(...), senha: str = Form(...)):
    res = (
        banco_dados.table("users")
        .select("id_user, nome, email, pontos, is_admin, senha_hash")
        .eq("email", email.lower().strip())
        .execute()
    )

    if not res.data:
        raise HTTPException(status_code=401, detail="E-mail ou senha incorretos.")

    usuario = res.data[0]

    if not usuario.get("senha_hash"):
        raise HTTPException(
            status_code=409,
            detail="Esta conta ainda não tem senha definida. Configure uma senha para continuar.",
        )

    if not verificar_senha(senha, usuario["senha_hash"]):
        raise HTTPException(status_code=401, detail="E-mail ou senha incorretos.")

    token = criar_token(usuario["id_user"])
    usuario.pop("senha_hash", None)
    return {"user": usuario, "token": token}

@app.post("/usuarios/definir-senha")
@limiter.limit("5/minute")
def definir_senha(
    request: Request,
    email: str = Form(...),
    data_nasc: str = Form(...),
    nova_senha: str = Form(...),
):
    """
    Fluxo de transição para contas criadas antes da senha existir: usa a data de
    nascimento (segredo antigo) como prova única para permitir configurar a primeira senha.
    """
    v_senha, m_senha = validar_senha(nova_senha)
    if not v_senha:
        raise HTTPException(status_code=400, detail=m_senha)

    res = (
        banco_dados.table("users")
        .select("id_user, nome, email, pontos, is_admin, senha_hash")
        .eq("email", email.lower().strip())
        .eq("data_nasc", data_nasc)
        .execute()
    )
    if not res.data:
        raise HTTPException(status_code=401, detail="E-mail ou data de nascimento incorretos.")

    usuario = res.data[0]
    if usuario.get("senha_hash"):
        raise HTTPException(status_code=400, detail="Esta conta já tem senha definida. Faça login normalmente.")

    banco_dados.table("users").update({"senha_hash": hash_senha(nova_senha)}).eq("id_user", usuario["id_user"]).execute()

    token = criar_token(usuario["id_user"])
    usuario.pop("senha_hash", None)
    return {"user": usuario, "token": token}

@app.post("/responder")
def responder_pergunta(
    id_pergunta: str = Form(...),
    resposta: str = Form(...),
    tempo_segundos: int = Form(...),
    id_user: str = Depends(get_current_user),
):
    pergunta = banco_dados.table("perguntas").select("*").eq("id_pergunta", id_pergunta).single().execute()
    if not pergunta.data:
        raise HTTPException(status_code=404, detail="Pergunta não encontrada.")

    qrs_com_pergunta = (
        banco_dados.table("qrcodes")
        .select("code_hash")
        .eq("id_pergunta", id_pergunta)
        .execute()
    )
    hashes_validos = [q["code_hash"] for q in (qrs_com_pergunta.data or [])]

    if not hashes_validos:
        raise HTTPException(status_code=400, detail="Nenhum QR Code vinculado a esta pergunta.")

    captura_valida = (
        banco_dados.table("catch")
        .select("id_catch")
        .eq("id_user", id_user)
        .in_("code_hash", hashes_validos)
        .execute()
    )
    if not captura_valida.data:
        raise HTTPException(status_code=403, detail="Capture o QR Code antes de responder.")

    ja_respondeu = (
        banco_dados.table("user_perguntas")
        .select("id")
        .eq("id_user", id_user)
        .eq("id_pergunta", id_pergunta)
        .execute()
    )
    if ja_respondeu.data:
        raise HTTPException(status_code=400, detail="Você já respondeu esta pergunta.")

    p = pergunta.data
    acertou = resposta.upper() == p["resposta_correta"].upper()
    pontos_bonus = (p["pontos_rapido"] if tempo_segundos <= 10 else p["pontos_lento"]) if acertou else 0

    try:
        # Sempre registra, independente de acerto — o trigger cuida dos pontos
        banco_dados.table("user_perguntas").insert({
            "id_user": id_user,
            "id_pergunta": id_pergunta,
            "resposta": resposta,
            "tempo_segundos": tempo_segundos,
        }).execute()
    except Exception:
        raise HTTPException(status_code=500, detail="Erro ao registrar resposta. Tente novamente.")

    return {
        "acertou": acertou,
        "resposta_correta": p["resposta_correta"],
        "pontos_bonus": pontos_bonus,
        "feedback": "Resposta correta! 🎉" if acertou else "Resposta errada. Sem pontos bônus desta vez.",
    }

@app.get("/usuarios/me/medalhas")
def medalhas_do_usuario(id_user: str = Depends(get_current_user)):
    response = (
        banco_dados.table("user_medalhas")
        .select("id_medalha, conquistado_em, medalhas(id_medalha, nome, descricao, imagem_base64)")
        .eq("id_user", id_user)
        .order("conquistado_em", desc=True)
        .execute()
    )
    if not response.data:
        return []
    resultado = []
    for row in response.data:
        medalha = row.get("medalhas")
        if medalha:
            resultado.append({
                "id_medalha": medalha["id_medalha"],
                "nome": medalha["nome"],
                "descricao": medalha.get("descricao", ""),
                "imagem_base64": medalha.get("imagem_base64", ""),
                "conquistado_em": row["conquistado_em"],
            })
    return resultado

@app.patch("/usuarios/me/nome")
def atualizar_nome(nome: str = Form(...), id_user: str = Depends(get_current_user)):
    v_nome, m_nome = validar_nome_sem_numeros(nome)
    if not v_nome:
        raise HTTPException(status_code=400, detail=m_nome)

    resultado = banco_dados.table("users").update({"nome": nome.title()}).eq("id_user", id_user).execute()
    if not resultado.data:
        raise HTTPException(status_code=404, detail="Usuário não encontrado.")
    return {"status": "Sucesso", "nome": resultado.data[0]["nome"]}

# ==================== QR PESSOAL / AMIGOS (XP) ====================

@app.get("/usuarios/me/qrcode")
def meu_qrcode(id_user: str = Depends(get_current_user)):
    usuario = banco_dados.table("users").select("personal_code_hash").eq("id_user", id_user).single().execute()
    if not usuario.data:
        raise HTTPException(status_code=404, detail="Usuário não encontrado.")

    personal_code_hash = usuario.data.get("personal_code_hash")
    if not personal_code_hash:
        personal_code_hash = gerar_code_hash(id_user)
        banco_dados.table("users").update({"personal_code_hash": personal_code_hash}).eq("id_user", id_user).execute()

    qr = qrcode.QRCode(version=1, box_size=10, border=5)
    qr.add_data(f"https://qr-code-hunt.vercel.app/perfil/{personal_code_hash}")
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return StreamingResponse(buf, media_type="image/png")

@app.post("/amigos/escanear")
def escanear_amigo(code_hash: str = Form(...), id_user: str = Depends(get_current_user)):
    alvo = banco_dados.table("users").select("id_user, nome").eq("personal_code_hash", code_hash).execute()
    if not alvo.data:
        raise HTTPException(status_code=404, detail="QR Code de amigo não encontrado.")

    scanned = alvo.data[0]
    scanned_id = scanned["id_user"]
    if scanned_id == id_user:
        raise HTTPException(status_code=400, detail="Você não pode escanear seu próprio QR Code.")

    # Networking é mútuo: bloqueia se essa dupla já se conectou em qualquer direção
    ja_conectados = (
        banco_dados.table("friend_scans")
        .select("id")
        .or_(
            f"and(scanner_id.eq.{id_user},scanned_id.eq.{scanned_id}),"
            f"and(scanner_id.eq.{scanned_id},scanned_id.eq.{id_user})"
        )
        .execute()
    )
    if ja_conectados.data:
        raise HTTPException(status_code=409, detail=f"Você e {scanned['nome']} já se conectaram antes.")

    try:
        banco_dados.table("friend_scans").insert({
            "scanner_id": id_user,
            "scanned_id": scanned_id,
        }).execute()
    except Exception as e:
        error_msg = str(e)
        if "23505" in error_msg or "duplicate key" in error_msg:
            raise HTTPException(status_code=409, detail=f"Você e {scanned['nome']} já se conectaram antes.")
        raise HTTPException(status_code=500, detail="Erro ao registrar a captura de amigo.")

    # Networking: os dois lados ganham pontos pela conexão
    scanner = banco_dados.table("users").select("pontos").eq("id_user", id_user).single().execute()
    pontos_atual = scanner.data.get("pontos") or 0 if scanner.data else 0
    pontos_novo = pontos_atual + PONTOS_POR_AMIGO
    banco_dados.table("users").update({"pontos": pontos_novo}).eq("id_user", id_user).execute()

    alvo_atual = banco_dados.table("users").select("pontos").eq("id_user", scanned_id).single().execute()
    alvo_pontos_atual = alvo_atual.data.get("pontos") or 0 if alvo_atual.data else 0
    banco_dados.table("users").update({"pontos": alvo_pontos_atual + PONTOS_POR_AMIGO}).eq("id_user", scanned_id).execute()

    nivel_anterior = calcular_nivel(pontos_atual)
    nivel_atual = calcular_nivel(pontos_novo)

    return {
        "status": "Sucesso",
        "amigo": scanned["nome"],
        "pontos_ganho": PONTOS_POR_AMIGO,
        "pontos_total": pontos_novo,
        "nivel_anterior": nivel_anterior,
        "nivel_atual": nivel_atual,
        "subiu_de_nivel": nivel_atual > nivel_anterior,
    }

# ==================== QRCODE (EVENTO) ====================

@app.post("/qrcodes/gerar")
def gerar_qr(
    nome_local: str,
    pontos: int,
    admin_id: str = Depends(require_admin),
    id_medalha: str = Query(None),
    id_pergunta: str = Query(None),

):
    code_hash = gerar_code_hash(f"{nome_local}-{pontos}")

    insert_data = {
        "code_hash": code_hash,
        "pontos": pontos,
        "local": nome_local,
    }
    if id_medalha:
        insert_data["id_medalha"] = id_medalha
    if id_pergunta:
        insert_data["id_pergunta"] = id_pergunta

    banco_dados.table('qrcodes').insert(insert_data).execute()

    return {"status": "Sucesso", "code_hash": code_hash}


@app.get("/qrcodes/download/{code_hash}")
def download_qr(code_hash: str):
    response = banco_dados.table("qrcodes").select("local").eq("code_hash", code_hash).execute()
    if not response.data:
        raise HTTPException(status_code=404, detail="QR Code não encontrado")

    local = response.data[0]["local"] or code_hash

    qr = qrcode.QRCode(version=1, box_size=10, border=5)
    qr.add_data(f"https://qr-code-hunt.vercel.app/scan/{code_hash}")
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)

    nome_arquivo = re.sub(r"[^a-zA-Z0-9_-]+", "_", local).strip("_") or code_hash
    return StreamingResponse(
        buf,
        media_type="image/png",
        headers={"Content-Disposition": f"attachment; filename=qrcode_{nome_arquivo}.png"},
    )

@app.get("/qrcodes/{code_hash}/pdf")
def download_qr_pdf(code_hash: str, admin_id: str = Depends(require_admin)):
    """Gera uma etiqueta pronta pra imprimir: QR Code + nome do local + pontuação."""
    response = banco_dados.table("qrcodes").select("local, pontos").eq("code_hash", code_hash).execute()
    if not response.data:
        raise HTTPException(status_code=404, detail="QR Code não encontrado")

    local = response.data[0]["local"] or "QR Code"
    pontos = response.data[0]["pontos"]

    qr = qrcode.QRCode(version=1, box_size=10, border=4)
    qr.add_data(f"https://qr-code-hunt.vercel.app/scan/{code_hash}")
    qr.make(fit=True)
    qr_img = qr.make_image(fill_color="black", back_color="white").convert("RGB")

    largura, altura = 600, 800
    etiqueta = Image.new("RGB", (largura, altura), "white")
    desenho = ImageDraw.Draw(etiqueta)

    fonte_titulo = ImageFont.load_default(size=32)
    fonte_legenda = ImageFont.load_default(size=24)

    def centralizar(texto, fonte, y):
        caixa = desenho.textbbox((0, 0), texto, font=fonte)
        x = (largura - (caixa[2] - caixa[0])) / 2
        desenho.text((x, y), texto, fill="black", font=fonte)

    centralizar(local, fonte_titulo, 40)

    qr_tamanho = 440
    qr_redimensionado = qr_img.resize((qr_tamanho, qr_tamanho))
    etiqueta.paste(qr_redimensionado, ((largura - qr_tamanho) // 2, 120))

    centralizar(f"+{pontos} pontos", fonte_legenda, 120 + qr_tamanho + 30)
    centralizar("Caça QR Code - UNDB", fonte_legenda, altura - 60)

    buf = io.BytesIO()
    etiqueta.save(buf, format="PDF")
    buf.seek(0)

    nome_arquivo = re.sub(r"[^a-zA-Z0-9_-]+", "_", local).strip("_") or code_hash
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=qrcode_{nome_arquivo}.pdf"},
    )

@app.get("/qrcodes/listar")
def listar_qrcodes(admin_id: str = Depends(require_admin)):
    response = banco_dados.table('qrcodes').select("*").execute()
    return response.data

class StatusUpdate(BaseModel):
    ativo: bool

@app.patch("/qrcodes/status/{code_hash}")
def toggle_status_qr(code_hash: str, body: StatusUpdate, admin_id: str = Depends(require_admin)):
    resultado = banco_dados.table('qrcodes').update({"ativo": body.ativo}).eq("code_hash", code_hash).execute()
    if not resultado.data:
        raise HTTPException(status_code=404, detail="QR Code não encontrado")
    return {"status": "sucesso", "ativo": body.ativo}

@app.patch("/qrcodes/{code_hash}/vincular")
def vincular_qrcode(
    code_hash: str,
    admin_id: str = Depends(require_admin),
    id_medalha: str = Query(None),
    id_pergunta: str = Query(None),
):
    atualizacao = {}
    if id_medalha is not None:
        atualizacao["id_medalha"] = None if id_medalha == "null" else id_medalha
    if id_pergunta is not None:
        atualizacao["id_pergunta"] = None if id_pergunta == "null" else id_pergunta

    if not atualizacao:
        raise HTTPException(status_code=400, detail="Nenhum campo para atualizar.")

    resultado = banco_dados.table("qrcodes").update(atualizacao).eq("code_hash", code_hash).execute()
    if not resultado.data:
        raise HTTPException(status_code=404, detail="QR Code não encontrado.")
    return {"status": "Sucesso", "qrcode": resultado.data[0]}


@app.post("/capturar")
def capturar(code_hash: str = Form(...), id_user: str = Depends(get_current_user)):

    try:

        qr_data = (
        banco_dados.table("qrcodes")
        .select("pontos, ativo, id_medalha, id_pergunta")
        .eq("code_hash", code_hash)
        .single()
        .execute()
        )
        if not qr_data.data:
            raise HTTPException(status_code=404, detail="QR Code não encontrado.")
        if not qr_data.data["ativo"]:
            raise HTTPException(status_code=403, detail="Este QR Code foi desativado.")
        # 1 Verifica se o QRcode existe ou se esta ativo


        ja_capturado = (
            banco_dados.table("catch")
            .select("id_catch")
            .eq("id_user", id_user)
            .eq("code_hash", code_hash)
            .execute()
        )
        # 2 verifica se ja foi capturado
        if ja_capturado.data:
            raise HTTPException(status_code=409, detail="Você já capturou este QR Code!")

        valor_pontos = qr_data.data["pontos"]
        id_medalha   = qr_data.data["id_medalha"]
        id_pergunta  = qr_data.data["id_pergunta"]

        antes = banco_dados.table("users").select("pontos").eq("id_user", id_user).single().execute()
        pontos_antes = (antes.data or {}).get("pontos") or 0

        # 3. Registra a captura do QR (o trigger trg_atualiza_pontos soma os pontos)
        banco_dados.table("catch").insert({
            "id_user": id_user,
            "catch_time": datetime.now().isoformat(),
            "code_hash": code_hash
        }).execute()

        depois = banco_dados.table("users").select("pontos").eq("id_user", id_user).single().execute()
        pontos_depois = (depois.data or {}).get("pontos") or 0
        nivel_anterior = calcular_nivel(pontos_antes)
        nivel_atual = calcular_nivel(pontos_depois)

        # 5. Tenta conceder a medalha — mas só se o usuário ainda não a tiver
        medalha_conquistada = False
        if id_medalha:
            ja_tem_medalha = (
                banco_dados.table("user_medalhas")
                .select("id")
                .eq("id_user", id_user)
                .eq("id_medalha", id_medalha)
                .execute()
            )
            if not ja_tem_medalha.data:
                banco_dados.table("user_medalhas").insert({
                    "id_user": id_user,
                    "id_medalha": id_medalha,
                    "conquistado_em": datetime.now().isoformat()
                }).execute()
                medalha_conquistada = True
            # Se já tinha a medalha, simplesmente ignora — o QR já foi capturado com sucesso acima

        # 6. Busca a pergunta vinculada, se houver
        pergunta = None
        if id_pergunta:
            p = banco_dados.table("perguntas").select("*").eq("id_pergunta", id_pergunta).single().execute()
            if p.data:
                pergunta = {
                    "id_pergunta": p.data["id_pergunta"],
                    "enunciado": p.data["enunciado"],
                    "tipo": p.data["tipo"],
                    "alternativas": p.data["alternativas"],
                }

        return {
            "status": "Sucesso",
            "pontos_qr": valor_pontos,
            "pontos_total": pontos_depois,
            "nivel_anterior": nivel_anterior,
            "nivel_atual": nivel_atual,
            "subiu_de_nivel": nivel_atual > nivel_anterior,
            "pergunta": pergunta,
            "medalha_conquistada": medalha_conquistada  # True só se a medalha foi de fato concedida agora
        }

    except HTTPException:
        raise
    except Exception as e:
        print(f"Erro inesperado na captura: {e}")
        raise HTTPException(status_code=500, detail="Erro interno ao processar a captura. Tente novamente.")

# ==================== MEDALHAS ====================

@app.post("/medalhas/nova")
async def criar_medalha(
    admin_id: str = Depends(require_admin),
    nome: str = Form(...),
    descricao: str = Form(""),
    imagem: UploadFile = File(...)
):
    conteudo = await imagem.read()
    base64_img = base64.b64encode(conteudo).decode("utf-8")
    mime = imagem.content_type
    imagem_base64 = f"data:{mime};base64,{base64_img}"

    resultado = banco_dados.table("medalhas").insert({
        "nome": nome,
        "descricao": descricao,
        "imagem_base64": imagem_base64
    }).execute()

    return {"status": "Sucesso", "medalha": resultado.data[0]}

@app.get("/medalhas/listar")
def listar_medalhas(admin_id: str = Depends(require_admin)):
    response = banco_dados.table("medalhas").select("*").order("criado_em", desc=True).execute()
    return response.data or []

@app.patch("/medalhas/{id_medalha}")
async def editar_medalha(
    id_medalha: str,
    admin_id: str = Depends(require_admin),
    nome: str = Form(None),
    descricao: str = Form(None),
    imagem: UploadFile = File(None),
):
    atualizacao = {}
    if nome is not None:
        atualizacao["nome"] = nome
    if descricao is not None:
        atualizacao["descricao"] = descricao
    if imagem and imagem.filename:
        conteudo = await imagem.read()
        base64_img = base64.b64encode(conteudo).decode("utf-8")
        mime = imagem.content_type
        atualizacao["imagem_base64"] = f"data:{mime};base64,{base64_img}"

    if not atualizacao:
        raise HTTPException(status_code=400, detail="Nenhum campo para atualizar.")

    resultado = banco_dados.table("medalhas").update(atualizacao).eq("id_medalha", id_medalha).execute()
    if not resultado.data:
        raise HTTPException(status_code=404, detail="Medalha não encontrada.")
    return {"status": "Sucesso", "medalha": resultado.data[0]}

@app.delete("/medalhas/{id_medalha}")
def deletar_medalha(id_medalha: str, admin_id: str = Depends(require_admin)):
    banco_dados.table("qrcodes").update({"id_medalha": None}).eq("id_medalha", id_medalha).execute()
    banco_dados.table("medalhas").delete().eq("id_medalha", id_medalha).execute()
    return {"status": "Sucesso", "mensagem": "Medalha removida."}

# ==================== PERGUNTAS ====================

@app.post("/perguntas/nova")
def criar_pergunta(
    admin_id: str = Depends(require_admin),
    enunciado: str = Form(...),
    tipo: str = Form(...),
    resposta_correta: str = Form(...),
    pontos_rapido: int = Form(50),
    pontos_lento: int = Form(20),
    alternativa_a: str = Form(""),
    alternativa_b: str = Form(""),
    alternativa_c: str = Form(""),
    alternativa_d: str = Form(""),
):
    if tipo not in ("multipla_escolha", "verdadeiro_falso"):
        raise HTTPException(status_code=400, detail="Tipo inválido.")

    if tipo == "multipla_escolha":
        if not all([alternativa_a, alternativa_b, alternativa_c, alternativa_d]):
            raise HTTPException(status_code=400, detail="Preencha todas as alternativas A, B, C e D.")
        if resposta_correta not in ("A", "B", "C", "D"):
            raise HTTPException(status_code=400, detail="Resposta correta deve ser A, B, C ou D.")
        alternativas = {"A": alternativa_a, "B": alternativa_b, "C": alternativa_c, "D": alternativa_d}
    else:
        if resposta_correta not in ("V", "F"):
            raise HTTPException(status_code=400, detail="Resposta correta deve ser V ou F.")
        alternativas = {"V": "Verdadeiro", "F": "Falso"}

    resultado = banco_dados.table("perguntas").insert({
        "enunciado": enunciado,
        "tipo": tipo,
        "alternativas": alternativas,
        "resposta_correta": resposta_correta,
        "pontos_rapido": pontos_rapido,
        "pontos_lento": pontos_lento,
    }).execute()

    return {"status": "Sucesso", "pergunta": resultado.data[0]}

@app.get("/perguntas/listar")
def listar_perguntas(admin_id: str = Depends(require_admin)):
    response = banco_dados.table("perguntas").select("*").order("criado_em", desc=True).execute()
    return response.data or []

@app.patch("/perguntas/{id_pergunta}")
def editar_pergunta(
    id_pergunta: str,
    admin_id: str = Depends(require_admin),
    enunciado: str = Form(None),
    tipo: str = Form(None),
    resposta_correta: str = Form(None),
    pontos_rapido: int = Form(None),
    pontos_lento: int = Form(None),
    alternativa_a: str = Form(None),
    alternativa_b: str = Form(None),
    alternativa_c: str = Form(None),
    alternativa_d: str = Form(None),
):
    atualizacao = {}
    if enunciado is not None:
        atualizacao["enunciado"] = enunciado
    if tipo is not None:
        atualizacao["tipo"] = tipo
    if resposta_correta is not None:
        atualizacao["resposta_correta"] = resposta_correta
    if pontos_rapido is not None:
        atualizacao["pontos_rapido"] = pontos_rapido
    if pontos_lento is not None:
        atualizacao["pontos_lento"] = pontos_lento

    # Determina o tipo final (pode ter mudado nesta requisição ou já existia)
    tipo_final = tipo
    if tipo_final is None:
        current = banco_dados.table("perguntas").select("tipo, alternativas").eq("id_pergunta", id_pergunta).single().execute()
        tipo_final = current.data.get("tipo") if current.data else None

    if tipo_final == "verdadeiro_falso":
        # Sempre sobrescreve com apenas V/F, descartando qualquer A/B/C/D anterior
        atualizacao["alternativas"] = {"V": "Verdadeiro", "F": "Falso"}
    elif tipo_final == "multipla_escolha":
        if any(x is not None for x in [alternativa_a, alternativa_b, alternativa_c, alternativa_d]):
            # Busca alternativas atuais SÓ se o tipo não mudou (para preservar as que não foram enviadas)
            if tipo is None:  # tipo não mudou, faz merge
                current = banco_dados.table("perguntas").select("alternativas").eq("id_pergunta", id_pergunta).single().execute()
                alts = current.data.get("alternativas", {}) if current.data else {}
                # Limpa chaves de verdadeiro_falso que possam ter ficado
                alts = {k: v for k, v in alts.items() if k in ("A", "B", "C", "D")}
            else:  # tipo mudou para multipla_escolha, começa do zero
                alts = {}
            if alternativa_a is not None: alts["A"] = alternativa_a
            if alternativa_b is not None: alts["B"] = alternativa_b
            if alternativa_c is not None: alts["C"] = alternativa_c
            if alternativa_d is not None: alts["D"] = alternativa_d
            atualizacao["alternativas"] = alts

    if not atualizacao:
        raise HTTPException(status_code=400, detail="Nenhum campo para atualizar.")

    resultado = banco_dados.table("perguntas").update(atualizacao).eq("id_pergunta", id_pergunta).execute()
    if not resultado.data:
        raise HTTPException(status_code=404, detail="Pergunta não encontrada.")
    return {"status": "Sucesso", "pergunta": resultado.data[0]}

@app.delete("/perguntas/{id_pergunta}")
def deletar_pergunta(id_pergunta: str, admin_id: str = Depends(require_admin)):
    banco_dados.table("qrcodes").update({"id_pergunta": None}).eq("id_pergunta", id_pergunta).execute()
    banco_dados.table("perguntas").delete().eq("id_pergunta", id_pergunta).execute()
    return {"status": "Sucesso", "mensagem": "Pergunta removida."}
